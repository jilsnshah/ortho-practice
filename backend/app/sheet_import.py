"""Read the practice's Excel ledger into the shape this app stores.

The sheet is one tab per city; inside a tab, a **bold** row names a clinic
("Avdhoot - Science City", "Dr. Harsh Verma - Althan"), the rows under it are
patients, and a row with no name but numbers is that clinic's hand-made
subtotal. Amounts repeat within a clinic, so the most common
quote/fee/material triple per treatment becomes the clinic's rate card.

Parsing lives here, apart from the CLI in ``scripts/import_sheet.py``, so it can
be tested without a workbook on disk or a server to talk to.
"""

from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from app.naming import name_key

HEADER_FIRST_CELL = "clinic / patient name"
# "Avdhoot - Science City" -> clinic "Avdhoot", branch "Science City".
BRANCH_SEPARATOR = " - "


@dataclass
class SheetRow:
    patient: str
    treatment: str | None
    quote_amount: int
    fee_amount: int
    material_cost: int
    ref: str
    warnings: list[str] = field(default_factory=list)

    @property
    def profit(self) -> int:
        return self.fee_amount - self.material_cost

    @property
    def is_priced(self) -> bool:
        return self.treatment is not None and self.fee_amount > 0


@dataclass
class ClinicBlock:
    name: str
    branch: str | None
    doctor: str | None
    ref: str
    rows: list[SheetRow] = field(default_factory=list)
    #: The sheet's own subtotal row, when it has one (quote, fee, material, profit).
    subtotal: tuple[int, int, int, int] | None = None

    @property
    def label(self) -> str:
        return f"{self.name} – {self.branch}" if self.branch else self.name

    def rate_card(self) -> dict[str, tuple[int, int, int]]:
        """The usual amounts per treatment at this clinic: the most common triple."""
        by_treatment: dict[str, Counter] = {}
        for row in self.rows:
            if not row.is_priced:
                continue
            counts = by_treatment.setdefault(row.treatment, Counter())
            counts[(row.quote_amount, row.fee_amount, row.material_cost)] += 1
        return {treatment: counts.most_common(1)[0][0] for treatment, counts in by_treatment.items()}

    def totals(self) -> tuple[int, int, int]:
        return (
            sum(r.quote_amount for r in self.rows),
            sum(r.fee_amount for r in self.rows),
            sum(r.material_cost for r in self.rows),
        )


@dataclass
class CityBlock:
    name: str
    clinics: list[ClinicBlock] = field(default_factory=list)


def _amount(value) -> tuple[int, bool]:
    """Money cell as whole rupees, plus whether it held a number at all."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0, False
    try:
        return int(round(float(value))), True
    except (TypeError, ValueError):
        return 0, False


def _split_clinic(title: str) -> tuple[str, str | None, str | None]:
    name, _, branch = title.partition(BRANCH_SEPARATOR)
    name, branch = name.strip(), branch.strip() or None
    # A clinic named after the doctor who runs it gives us the doctor too.
    doctor = name if name.lower().startswith(("dr.", "dr ")) else None
    return name, branch, doctor


def parse_workbook(path: str | Path) -> list[CityBlock]:
    from openpyxl import load_workbook  # imported lazily: only the import script needs it

    workbook = load_workbook(path, data_only=True)
    cities: list[CityBlock] = []

    for sheet in workbook.worksheets:
        city = CityBlock(name=sheet.title.strip())
        clinic: ClinicBlock | None = None

        for index, row in enumerate(sheet.iter_rows(values_only=False), start=1):
            cell = row[0]
            first = str(cell.value).strip() if cell.value is not None else ""
            values = [c.value for c in row]
            treatment = str(values[1]).strip() if len(values) > 1 and values[1] is not None else ""
            quote, has_quote = _amount(values[2] if len(values) > 2 else None)
            fee, has_fee = _amount(values[3] if len(values) > 3 else None)
            material, has_material = _amount(values[4] if len(values) > 4 else None)
            sheet_profit, _ = _amount(values[5] if len(values) > 5 else None)
            ref = f"{sheet.title.strip()}!A{index}"

            if first.lower() == HEADER_FIRST_CELL:
                continue

            if not first:
                # A nameless row carrying numbers is the clinic's own subtotal.
                if clinic is not None and (has_quote or has_fee or has_material):
                    clinic.subtotal = (quote, fee, material, sheet_profit)
                continue

            # Clinic headings are bold; patients never are.
            if cell.font is not None and cell.font.bold:
                name, branch, doctor = _split_clinic(first)
                clinic = ClinicBlock(name=name, branch=branch, doctor=doctor, ref=ref)
                city.clinics.append(clinic)
                continue

            if clinic is None:
                continue  # a stray row above the first clinic

            warnings = []
            if not treatment:
                warnings.append("no treatment recorded")
            if not has_fee:
                warnings.append("no fee recorded")
            if not has_quote and (has_fee or treatment):
                warnings.append("no patient quote recorded")
            if has_material and has_fee and material > fee:
                warnings.append(f"material ({material}) is more than the fee ({fee}) — negative profit")

            clinic.rows.append(
                SheetRow(
                    patient=first,
                    treatment=treatment or None,
                    quote_amount=quote,
                    fee_amount=fee,
                    material_cost=material,
                    ref=ref,
                    warnings=warnings,
                )
            )

        if city.clinics:
            cities.append(city)

    return cities


def treatment_spellings(cities: list[CityBlock]) -> dict[str, str]:
    """Map each treatment's normalized key to its most common spelling.

    "Self - Ligating Metal Braces", "Self-Ligating Metal Braces" and
    "self - Ligating Metal Braces" are one treatment type with one name.
    """
    counts: dict[str, Counter] = {}
    for city in cities:
        for clinic in city.clinics:
            for row in clinic.rows:
                if row.treatment:
                    counts.setdefault(name_key(row.treatment), Counter())[row.treatment] += 1
    return {key: spellings.most_common(1)[0][0] for key, spellings in counts.items()}


def subtotal_mismatches(cities: list[CityBlock]) -> list[str]:
    """Clinics where our row sums disagree with the sheet's own subtotal row.

    Includes the sheet's hand-typed profit, which is the figure this app replaces
    with fee - material, so any arithmetic slip in the sheet shows up here.
    """
    problems = []
    for city in cities:
        for clinic in city.clinics:
            if clinic.subtotal is None:
                continue
            quote, fee, material = clinic.totals()
            ours = (quote, fee, material, fee - material)
            for label, a, b in zip(("quote", "fee", "material", "profit"), ours, clinic.subtotal):
                # A missing subtotal cell reads as 0; only compare what the sheet filled in.
                if b and a != b:
                    problems.append(f"{city.name} · {clinic.label}: {label} sums to {a:,}, sheet says {b:,}")
    return problems
